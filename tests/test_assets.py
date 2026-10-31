import json
import os

from stages import assets


def test_graph_drawn_and_placeholders_swapped(tmp_path, monkeypatch):
    (tmp_path / "blog.md").write_text("## 1. A\n\n<!-- IMAGE: a -->\n\n### Parent Review\n\n<!-- GRAPH: a -->\n\n<!-- IMAGE: ghost -->\n")
    (tmp_path / "assets.json").write_text(json.dumps({"schools": [
        {"name": "A", "slug": "a", "image_url": None, "reviews": {"Academics": [80, 10], "Fees": [55, 30]}}]}))
    problems = assets.process(str(tmp_path))
    md = (tmp_path / "blog.md").read_text()
    assert "![Parent sentiment breakdown for A](graphs/a.png)" in md
    assert "<!--" not in md  # unknown and image-less placeholders are removed
    assert os.path.exists(tmp_path / "graphs" / "a.png")
    assert any("no building photo" in p for p in problems)
    assert "Pipeline: images and graphs" in (tmp_path / "notes.md").read_text()


def test_bad_review_data_gives_no_graph(tmp_path):
    (tmp_path / "blog.md").write_text("<!-- GRAPH: a -->\n")
    (tmp_path / "assets.json").write_text(json.dumps({"schools": [
        {"name": "A", "slug": "a", "image_url": None, "reviews": {"Academics": "lots"}}]}))
    problems = assets.process(str(tmp_path))
    assert any("no category review percentages" in p for p in problems)
    assert "Pipeline: images and graphs" in (tmp_path / "notes.md").read_text()
    assert "<!--" not in (tmp_path / "blog.md").read_text()


def test_docx_export(tmp_path):
    from docx import Document
    from stages import docx_export
    (tmp_path / "blog.md").write_text(
        "Intro about schools in Kalkaji.\n\n## 1. A School\n\n[btn_admissions id=\"1\"]\n\n"
        "- [A School](https://yellowslate.com/delhi/a) is **great**.\n- two\n- three\n\n"
        "| Detail | Information |\n| --- | --- |\n| Board | CBSE |\n", encoding="utf-8")
    (tmp_path / "meta.json").write_text(json.dumps({"title": "T", "meta_title": "M", "meta_description": "D", "slug": "s"}))
    d = Document(docx_export.export(str(tmp_path)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Meta title: M" in text and '[btn_admissions id="1"]' in text and "A School is great." in text
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "CBSE"


def test_all_zero_reviews_draw_no_graph(tmp_path):
    """The Kolar run published nine schools with every category at 0/0, which the old
    valid_reviews() accepted because it only type-checked. The chart script drew a titled,
    axis-labelled, completely empty frame and QC passed it, because QC only asks whether a
    graph image exists. Zero is the absence of a measurement, not a measurement of zero.
    """
    (tmp_path / "blog.md").write_text("<!-- GRAPH: a -->\n")
    (tmp_path / "assets.json").write_text(json.dumps({"schools": [
        {"name": "A", "slug": "a", "image_url": None,
         "reviews": {"Academics": [0, 0], "Infrastructure": [0, 0]}}]}))
    problems = assets.process(str(tmp_path))
    assert not os.path.exists(tmp_path / "graphs" / "a.png")
    assert any("every category is 0" in p for p in problems)
    assert "<!--" not in (tmp_path / "blog.md").read_text()


def test_one_zero_category_still_draws_a_graph(tmp_path):
    """Only an entirely empty set is rejected. A single category with no signal sits beside
    real ones all the time, and dropping the whole graph for it would lose real data."""
    (tmp_path / "blog.md").write_text("<!-- GRAPH: a -->\n")
    (tmp_path / "assets.json").write_text(json.dumps({"schools": [
        {"name": "A", "slug": "a", "image_url": None,
         "reviews": {"Academics": [82, 11], "Fee Satisfaction": [0, 0]}}]}))
    problems = assets.process(str(tmp_path))
    assert os.path.exists(tmp_path / "graphs" / "a.png")
    assert not any("every category is 0" in p for p in problems)


def test_an_oversized_photo_is_shrunk_to_something_their_endpoint_accepts(tmp_path, monkeypatch):
    # Oxford's campus shot arrived at 5120x2880 and 5.7 MB, over their 5 MB media limit, so
    # the draft carried a local path where the photo should have been.
    from PIL import Image
    monkeypatch.setattr(assets, "MAX_SIDE", 200)
    monkeypatch.setattr(assets, "MAX_BYTES", 20_000)
    path = tmp_path / "big.png"
    Image.frombytes("RGB", (1200, 800), os.urandom(1200 * 800 * 3)).save(path)
    assert path.stat().st_size > 20_000

    out = assets.shrink(str(path))

    assert os.path.getsize(out) <= 20_000
    with Image.open(out) as im:
        assert max(im.size) <= 200
