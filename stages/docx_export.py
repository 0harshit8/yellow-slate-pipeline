"""blog.md (+ meta.json, images, graphs) -> blog.docx, for sharing with Yellow Slate.

markdown-it parses the markdown, python-docx builds the document. Shortcodes like
[btn_admissions id="..."] are kept as visible grey text: they have to be pasted into
the CMS as-is. notes.md is internal and is not included.
"""

import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

GREY = RGBColor(0x80, 0x80, 0x80)
LINK = RGBColor(0x1A, 0x5F, 0xB4)
SHORTCODE = re.compile(r"(\[(?:btn_admissions|read_more_blogs|yellowslate_reviews|sponsored)\b[^\]]*\])")


def _hyperlink(paragraph, url, text, bold=False, italic=False):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A5FB4")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if bold:
        props.append(OxmlElement("w:b"))
    if italic:
        props.append(OxmlElement("w:i"))
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _text(paragraph, text, bold=False, italic=False):
    """Plain text, with shortcodes split out in grey."""
    for piece in SHORTCODE.split(text):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        run.bold, run.italic = bold, italic
        if SHORTCODE.fullmatch(piece):
            run.font.color.rgb = GREY
            run.font.size = Pt(9)


def _inline(paragraph, token, folder, doc):
    """Render an inline token's children into a paragraph."""
    bold = italic = False
    href = None
    link_text = []
    for child in token.children or []:
        t = child.type
        if t == "strong_open":
            bold = True
        elif t == "strong_close":
            bold = False
        elif t == "em_open":
            italic = True
        elif t == "em_close":
            italic = False
        elif t == "link_open":
            href, link_text = child.attrs.get("href"), []
        elif t == "link_close":
            text = "".join(link_text)
            if href and href.startswith("http"):
                _hyperlink(paragraph, href, text, bold, italic)
            else:
                _text(paragraph, text, bold, italic)
            href = None
        elif t in ("text", "code_inline"):
            if href is not None:
                link_text.append(child.content)
            else:
                _text(paragraph, child.content, bold, italic)
        elif t in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break() if t == "hardbreak" else paragraph.add_run(" ")
        elif t == "image":
            src = child.attrs.get("src", "")
            path = os.path.join(folder, src)
            if os.path.exists(path):
                width = Inches(6.0) if "graphs/" in src else Inches(5.0)
                try:
                    paragraph.add_run().add_picture(path, width=width)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    _text(paragraph, f"[image could not be embedded: {src}]")
            else:
                _text(paragraph, f"[missing image: {src}]")


def _table(doc, tokens, i, folder):
    """Build a table from tokens starting at table_open. Returns index after table_close."""
    rows, row, cell_inline = [], None, None
    while tokens[i].type != "table_close":
        t = tokens[i]
        if t.type == "tr_open":
            row = []
        elif t.type == "tr_close":
            rows.append(row)
        elif t.type in ("th_open", "td_open"):
            cell_inline = None
        elif t.type == "inline":
            cell_inline = t
        elif t.type in ("th_close", "td_close"):
            row.append((cell_inline, t.type == "th_close"))
        i += 1
    if rows:
        table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        table.style = "Table Grid"
        for r, cells in enumerate(rows):
            for c, (inline, header) in enumerate(cells):
                p = table.cell(r, c).paragraphs[0]
                if inline is not None:
                    _inline(p, inline, folder, doc)
                if header:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()
    return i + 1


def export(folder):
    """Writes folder/blog.docx and returns its path."""
    with open(os.path.join(folder, "blog.md"), encoding="utf-8") as fh:
        md = fh.read()
    meta_path = os.path.join(folder, "meta.json")
    meta = json.load(open(meta_path, encoding="utf-8")) if os.path.exists(meta_path) else {}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    if meta:
        doc.add_heading(meta.get("title", ""), level=0)
        for label, key in (("Meta title", "meta_title"), ("Meta description", "meta_description"), ("Slug", "slug")):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(str(meta.get(key, "")))
        doc.add_paragraph("_" * 60)

    tokens = MarkdownIt("commonmark").enable("table").parse(md)
    i, list_depth = 0, 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            level = int(t.tag[1])
            h = doc.add_heading(level=min(level, 4))
            _inline(h, tokens[i + 1], folder, doc)
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
        elif t.type == "paragraph_open":
            style = "List Bullet" if list_depth else None
            p = doc.add_paragraph(style=style)
            _inline(p, tokens[i + 1], folder, doc)
            i += 3
            continue
        elif t.type == "table_open":
            i = _table(doc, tokens, i, folder)
            continue
        elif t.type == "hr":
            doc.add_paragraph("_" * 60)
        i += 1

    out = os.path.join(folder, "blog.docx")
    try:
        doc.save(out)
    except PermissionError:
        raise PermissionError(f"{out} is open in Word. Close it and run: python pipeline.py docx")
    return out
